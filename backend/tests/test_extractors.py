"""
Unit test suite for FindIT multi-format extractors:
TextExtractor, PDFExtractor, DocxExtractor, PptxExtractor, and Extractor Factory.
"""

import tempfile
from pathlib import Path
import pytest
try:
    import pymupdf as fitz
except ImportError:
    import fitz
import docx
from pptx import Presentation

from app.extractors.text_extractor import TextExtractor
from app.extractors.pdf_extractor import PDFExtractor
from app.extractors.docx_extractor import DocxExtractor
from app.extractors.pptx_extractor import PptxExtractor
from app.extractors.factory import get_extractor_for_file


def test_text_extractor_plain_text():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "notes.txt"
        file_path.write_text("First paragraph.\nStill first.\n\nSecond paragraph.\nLine four.", encoding="utf-8")

        extractor = TextExtractor()
        doc = extractor.extract(file_path)

        assert doc.extraction_status == "OK"
        assert doc.file_type == "txt"
        assert len(doc.blocks) == 2

        # Block 1: lines 1-2
        assert doc.blocks[0].start_line == 1
        assert doc.blocks[0].end_line == 2
        assert "First paragraph." in doc.blocks[0].text
        assert doc.blocks[0].location_value == "Lines 1–2"

        # Block 2: lines 4-5
        assert doc.blocks[1].start_line == 4
        assert doc.blocks[1].end_line == 5
        assert "Second paragraph." in doc.blocks[1].text


def test_text_extractor_markdown_with_headings():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "doc.md"
        content = """# System Architecture
Overview of the system.

## Hybrid Storage
FAISS handles dense retrieval.
BM25 handles lexical search.

## Authentication
JWT token flow.
"""
        file_path.write_text(content, encoding="utf-8")

        extractor = TextExtractor()
        doc = extractor.extract(file_path)

        assert doc.extraction_status == "OK"
        assert doc.file_type == "md"
        assert len(doc.blocks) >= 2

        # Verify heading propagation
        headings = [b.section_heading for b in doc.blocks if b.section_heading]
        assert any("System Architecture" in h for h in headings)
        assert any("Hybrid Storage" in h for h in headings)


def test_text_extractor_code():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "auth.py"
        code = """import os

def verify_token(token: str) -> bool:
    # Validate token
    return bool(token)

class AuthService:
    def __init__(self):
        self.secret = 'test'
"""
        file_path.write_text(code, encoding="utf-8")

        extractor = TextExtractor()
        doc = extractor.extract(file_path)

        assert doc.extraction_status == "OK"
        assert doc.file_type == "code"
        assert len(doc.blocks) >= 1
        assert all(b.start_line is not None and b.end_line is not None for b in doc.blocks)
        assert any(b.section_heading == "verify_token" or b.section_heading == "AuthService" for b in doc.blocks)


def test_text_extractor_empty_file():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "empty.txt"
        file_path.write_text("   \n\n  ", encoding="utf-8")

        extractor = TextExtractor()
        doc = extractor.extract(file_path)

        assert doc.extraction_status == "NO_TEXT"
        assert len(doc.blocks) == 0


def test_pdf_extractor():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "sample.pdf"

        # Generate a 2-page test PDF with PyMuPDF
        pdf_doc = fitz.open()
        page1 = pdf_doc.new_page()
        page1.insert_text((50, 72), "This is Page 1 content for FindIT architecture spec.")
        page2 = pdf_doc.new_page()
        page2.insert_text((50, 72), "This is Page 2 content detailing vector indexing.")
        pdf_doc.save(file_path)
        pdf_doc.close()

        extractor = PDFExtractor()
        doc = extractor.extract(file_path)

        assert doc.extraction_status == "OK"
        assert doc.file_type == "pdf"
        assert doc.total_pages == 2
        assert len(doc.blocks) == 2

        # Page 1 check
        assert doc.blocks[0].page_number == 1
        assert doc.blocks[0].location_value == "Page 1 of 2"
        assert "Page 1 content" in doc.blocks[0].text

        # Page 2 check
        assert doc.blocks[1].page_number == 2
        assert doc.blocks[1].location_value == "Page 2 of 2"
        assert "Page 2 content" in doc.blocks[1].text


def test_pdf_extractor_scanned_no_text():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "blank.pdf"
        pdf_doc = fitz.open()
        pdf_doc.new_page()  # Page with no text
        pdf_doc.save(file_path)
        pdf_doc.close()

        extractor = PDFExtractor()
        doc = extractor.extract(file_path)

        assert doc.extraction_status == "NO_TEXT"
        assert len(doc.blocks) == 0


def test_docx_extractor():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "sample.docx"

        # Generate test DOCX with python-docx
        d = docx.Document()
        d.add_heading("Architecture Overview", level=1)
        d.add_paragraph("FindIT uses hybrid search.")
        d.add_paragraph("Second paragraph under architecture.")
        d.add_heading("Database Design", level=1)
        d.add_paragraph("SQLite WAL mode is enabled.")

        # Add table
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Key"
        t.cell(0, 1).text = "Value"
        t.cell(1, 0).text = "Engine"
        t.cell(1, 1).text = "FAISS"

        d.save(file_path)

        extractor = DocxExtractor()
        doc = extractor.extract(file_path)

        assert doc.extraction_status == "OK"
        assert doc.file_type == "docx"
        assert len(doc.blocks) >= 4

        headings = [b.section_heading for b in doc.blocks if b.section_heading]
        assert "Architecture Overview" in headings
        assert "Database Design" in headings


def test_pptx_extractor():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "sample.pptx"

        # Generate test PPTX with python-pptx
        prs = Presentation()
        blank_slide_layout = prs.slide_layouts[6]  # blank layout

        # Slide 1
        slide1 = prs.slides.add_slide(blank_slide_layout)
        tx1 = slide1.shapes.add_textbox(100, 100, 400, 100)
        tx1.text_frame.text = "Slide 1 Content: Search Pipeline"
        # Speaker notes
        notes1 = slide1.notes_slide
        notes1.notes_text_frame.text = "Remember to explain RRF formula here."

        # Slide 2
        slide2 = prs.slides.add_slide(blank_slide_layout)
        tx2 = slide2.shapes.add_textbox(100, 100, 400, 100)
        tx2.text_frame.text = "Slide 2 Content: Performance Benchmarks"

        prs.save(file_path)

        extractor = PptxExtractor()
        doc = extractor.extract(file_path)

        assert doc.extraction_status == "OK"
        assert doc.file_type == "pptx"
        assert doc.total_slides == 2
        assert len(doc.blocks) == 2

        # Check Slide 1
        assert doc.blocks[0].slide_number == 1
        assert "Search Pipeline" in doc.blocks[0].text
        assert "Remember to explain RRF formula" in doc.blocks[0].text
        assert "Slide 1" in doc.blocks[0].location_value

        # Check Slide 2
        assert doc.blocks[1].slide_number == 2
        assert "Performance Benchmarks" in doc.blocks[1].text
        assert "Slide 2" in doc.blocks[1].location_value


def test_extractor_factory():
    assert isinstance(get_extractor_for_file("doc.pdf"), PDFExtractor)
    assert isinstance(get_extractor_for_file("notes.docx"), DocxExtractor)
    assert isinstance(get_extractor_for_file("deck.pptx"), PptxExtractor)
    assert isinstance(get_extractor_for_file("readme.md"), TextExtractor)
    assert isinstance(get_extractor_for_file("main.py"), TextExtractor)
    assert isinstance(get_extractor_for_file("config.json"), TextExtractor)
    assert get_extractor_for_file("program.exe") is None
