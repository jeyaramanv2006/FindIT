"""
PptxExtractor utilizing python-pptx.
Extracts slide shapes, titles, and speaker notes preserving slide coordinates.
"""

from pathlib import Path
from typing import List, Optional
from app.extractors.base import BaseExtractor, ExtractedBlock, ExtractedDocument


class PptxExtractor(BaseExtractor):
    """
    Extracts text from PPTX presentations using python-pptx.
    Preserves exact slide numbers (1-indexed), slide titles, and speaker notes.
    """

    SUPPORTED_EXTS = {".pptx"}

    @classmethod
    def supported_extensions(cls) -> set[str]:
        return cls.SUPPORTED_EXTS

    def extract(self, file_path: Path) -> ExtractedDocument:
        target = Path(file_path).resolve()
        canonical_str = str(target).replace("\\", "/")
        ext = target.suffix.lower()

        if not target.exists():
            return ExtractedDocument(
                file_path=canonical_str,
                filename=target.name,
                extension=ext,
                file_type="pptx",
                extraction_status="ERROR",
                error_message=f"File not found: {target}",
            )

        try:
            from pptx import Presentation
        except ImportError:
            return ExtractedDocument(
                file_path=canonical_str,
                filename=target.name,
                extension=ext,
                file_type="pptx",
                extraction_status="ERROR",
                error_message="python-pptx is not installed.",
            )

        try:
            prs = Presentation(target)
        except Exception as e:
            return ExtractedDocument(
                file_path=canonical_str,
                filename=target.name,
                extension=ext,
                file_type="pptx",
                extraction_status="ERROR",
                error_message=f"Failed to open PPTX presentation: {e}",
            )

        total_slides = len(prs.slides)
        if total_slides == 0:
            return ExtractedDocument(
                file_path=canonical_str,
                filename=target.name,
                extension=ext,
                file_type="pptx",
                blocks=[],
                total_slides=0,
                extraction_status="NO_TEXT",
            )

        blocks: List[ExtractedBlock] = []

        for slide_idx, slide in enumerate(prs.slides):
            slide_number = slide_idx + 1
            slide_title: Optional[str] = None
            slide_texts: List[str] = []

            # 1. Inspect title shape
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()

            # 2. Extract shape text
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text and text != slide_title:
                        slide_texts.append(text)

            # 3. Extract speaker notes
            notes_text: Optional[str] = None
            if hasattr(slide, "notes_slide") and slide.notes_slide:
                notes_tf = getattr(slide.notes_slide, "notes_text_frame", None)
                if notes_tf and notes_tf.text.strip():
                    notes_text = notes_tf.text.strip()

            # Combine slide body
            combined_parts = []
            if slide_title:
                combined_parts.append(f"Title: {slide_title}")
            if slide_texts:
                combined_parts.extend(slide_texts)
            if notes_text:
                combined_parts.append(f"Notes: {notes_text}")

            if combined_parts:
                slide_content = "\n\n".join(combined_parts)
                loc_val = f"Slide {slide_number}"
                if slide_title:
                    loc_val += f" · {slide_title}"

                blocks.append(
                    ExtractedBlock(
                        text=slide_content,
                        location_type="slide",
                        location_value=loc_val,
                        slide_number=slide_number,
                        section_heading=slide_title,
                    )
                )

        status = "OK" if blocks else "NO_TEXT"

        return ExtractedDocument(
            file_path=canonical_str,
            filename=target.name,
            extension=ext,
            file_type="pptx",
            blocks=blocks,
            total_slides=total_slides,
            extraction_status=status,
        )
